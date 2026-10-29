"""Ingestion, retrieval, agent loop, workflows and exports (all offline)."""
from __future__ import annotations

import io
import json

from conftest import FakeLLM, make_pdf, message, tool_call

from clinstat.agent import run_agent
from clinstat.chunking import chunk_pages, split_sentences
from clinstat.embeddings import HashingEmbedder
from clinstat.export import record_to_docx, record_to_markdown
from clinstat.index_store import KnowledgeIndex, ingest, remove_document
from clinstat.pdf_text import clean_pages, strip_line_numbers
from clinstat.retrieval import BM25, tokenize
from clinstat.tools import normalize_citations
from clinstat.workflows import check_section, compare_documents, load_checklists


# ------------------------------------------------------------------ text processing
def test_line_numbers_stripped_and_headings_found():
    lines = ["Intro text that is long enough to count 12", "more words follow here 13", "and here 14",
             "even more words 15", "B. Strategies 16", "final words 17"]
    out = strip_line_numbers(lines)
    assert out[0].endswith("count") and out[4] == "B. Strategies"
    pages = clean_pages(["\n".join(lines)])
    assert any(b.kind == "heading" and b.text == "B. Strategies" for b in pages[0].blocks)


def test_reference_and_toc_pages_dropped():
    toc = "\n".join(f"{i}. Section title number {i} .................. {i * 3}" for i in range(1, 12))
    refs = "\n".join(
        f"{name}, A. B. (20{10 + i}). A paper on crossover trials. Stat Med {i + 1}:{i}0-{i}9."
        for i, name in enumerate(["Smith", "Jones", "Brown", "Senn", "Grizzle", "Hills", "Kenward", "Roger", "Lee", "Chen"])
    )
    body = "This is ordinary prose about cross-over trials. " * 10
    kinds = [p.kind for p in clean_pages([toc, body, refs])]
    assert kinds == ["toc", "content", "references"]


def test_hyphenation_rejoined_and_sentences_split():
    pages = clean_pages(["The analy-\nsis of cross-over trials is described here in some considerable detail. Next."])
    text = pages[0].blocks[0].text
    assert "analysis" in text
    assert split_sentences("See Fig. 2 and e.g. Table 3. Then it ends. Another one.") == [
        "See Fig. 2 and e.g. Table 3.", "Then it ends.", "Another one."]


def test_chunks_respect_size_and_overlap():
    long_page = " ".join(f"Sentence number {i} talks about estimands and intercurrent events." for i in range(120))
    chunks = chunk_pages(clean_pages([long_page]), chunk_chars=600, overlap_chars=120)
    assert len(chunks) > 5
    assert all(len(c.text) <= 600 for c in chunks)
    first_tail = chunks[0].text.split(". ")[-1]
    assert first_tail.rstrip(".") in chunks[1].text  # overlap carried forward


def test_tokenizer_handles_hyphens_and_plurals():
    toks = tokenize("Cross-over estimands, E9(R1) and PROC GLIMMIX")
    assert {"cross", "over", "crossover", "estimand", "e9", "r1", "proc", "glimmix"} <= set(toks)


def test_bm25_prefers_exact_term():
    bm = BM25(["the grizzle procedure for carryover", "estimands and intercurrent events", "carryover in general"])
    s = bm.scores("Grizzle carryover")
    assert s.argmax() == 0


# ------------------------------------------------------------------ ingestion
def test_ingest_report_and_incremental(built, kb):
    cfg, index, report = built
    assert sorted(report["added"]) == ["RegulatoryDocs/Estimand_Guidance_2020.pdf", "StatBooks/Crossover_Book_2015.pdf"]
    assert [s["doc_id"] for s in report["skipped"]] == ["StatPapers/scanned.pdf"]
    assert "scanned" in report["skipped"][0]["reason"]
    assert not any("ignored" in c["source"] for c in index.chunks)  # _Removed/ folder ignored
    book = index.documents["StatBooks/Crossover_Book_2015.pdf"]
    assert book["skipped_pages"].get("references") == 1
    assert index.embeddings.shape[0] == len(index.chunks)

    # Unchanged run does nothing; skipped file is not retried.
    again = ingest(cfg, HashingEmbedder(), log=lambda m: None)
    assert again["added"] == again["updated"] == again["skipped"] == [] and again["unchanged"] == 3

    # Modify one file, delete another.
    make_pdf(kb / "StatBooks" / "Crossover_Book_2015.pdf", [BOOK_EXTRA])
    (kb / "RegulatoryDocs" / "Estimand_Guidance_2020.pdf").unlink()
    third = ingest(cfg, HashingEmbedder(), log=lambda m: None)
    assert third["updated"] == ["StatBooks/Crossover_Book_2015.pdf"]
    assert third["removed"] == ["RegulatoryDocs/Estimand_Guidance_2020.pdf"]
    idx = KnowledgeIndex.load(cfg.index_dir)
    assert {c["doc_id"] for c in idx.chunks} == {"StatBooks/Crossover_Book_2015.pdf"}
    assert any("Kenward-Roger" in c["text"] for c in idx.chunks)


BOOK_EXTRA = ["3.1 Mixed models", "The Kenward-Roger adjustment corrects small-sample bias in standard errors of",
              "fixed effects estimated by REML in mixed models used for cross-over data."]


def test_remove_document(built):
    cfg, index, _ = built
    assert remove_document(cfg.index_dir, "StatBooks/Crossover_Book_2015.pdf")
    idx = KnowledgeIndex.load(cfg.index_dir)
    assert "StatBooks/Crossover_Book_2015.pdf" not in idx.documents
    assert all(c["doc_id"] != "StatBooks/Crossover_Book_2015.pdf" for c in idx.chunks)


def test_page_metadata_and_sections(built):
    _, index, _ = built
    strat = [c for c in index.chunks if "treatment policy strategy" in c["text"]]
    assert strat and strat[0]["page"] == 2
    assert "Strategies" in strat[0]["section"]
    assert not any(c["text"].rstrip().endswith((" 8", " 9")) for c in index.chunks)  # line numbers gone


# ------------------------------------------------------------------ retrieval
def test_hybrid_search_and_filters(retriever):
    hits = retriever.search("Grizzle two-stage carryover test", top_k=3)
    assert hits[0].chunk["source"] == "Crossover_Book_2015.pdf"
    only_reg = retriever.search("carryover", top_k=3, categories=["RegulatoryDocs"])
    assert all(h.chunk["category"] == "RegulatoryDocs" for h in only_reg)
    by_doc = retriever.search("analysis", top_k=3, sources=["Estimand_Guidance"])
    assert by_doc and all(h.chunk["source"] == "Estimand_Guidance_2020.pdf" for h in by_doc)


# ------------------------------------------------------------------ citations
def test_normalize_citations():
    text, cited, invalid = normalize_citations("A [S1, S3] B [S2; 4] C [C1] D [S9]", {"S1", "S2", "S3", "S4", "C1"})
    assert text == "A [S1][S3] B [S2][S4] C [C1] D "
    assert cited == ["S1", "S3", "S2", "S4", "C1"] and invalid == ["S9"]


# ------------------------------------------------------------------ agent loop
def test_agent_searches_calculates_and_cites(retriever):
    llm = FakeLLM([
        message(tool_calls=[
            tool_call("t1", "search_knowledge_bank", {"query": "Grizzle two-stage procedure"}),
            tool_call("t2", "continuous_sample_size", {"design": "crossover_2x2", "hypothesis": "superiority",
                                                       "delta": 0.5, "sd": 1.0}),
        ]),
        message(content="The two-stage procedure is not recommended [S1]. You need 66 subjects [C1]. Bogus [S42]."),
    ])
    result = run_agent("Should I use Grizzle's test? Sample size for d=0.5?", retriever, llm,
                       history=[{"role": "user", "content": "hi"}, {"role": "assistant", "content": "hello"}])
    assert result.error is None
    assert result.cited == ["S1", "C1"] and result.invalid_tags == ["S42"]
    assert "[S42]" not in result.answer
    assert result.citations.calcs[0].result.outputs["n_total"] == 66
    # second request carries the tool results and the history
    sent = llm.calls[1]["messages"]
    assert sent[1]["content"] == "hi" and sent[-1]["role"] == "tool"
    assert any(m["role"] == "tool" and "[S1]" in m["content"] for m in sent)
    rec = result.to_record()
    assert rec["sources"][0]["cited"] is True and json.dumps(rec)
    t = rec["timings"]
    assert {"total_s", "model_s", "search_s", "rerank_s", "tools_s"} <= set(t)
    assert t["total_s"] >= t["model_s"] and t["total_s"] >= t["search_s"]
    assert all("seconds" in e for e in result.tool_log)


def test_agent_forces_answer_after_tool_budget(retriever):
    loop = [message(tool_calls=[tool_call(f"t{i}", "search_knowledge_bank", {"query": "washout"})]) for i in range(4)]
    llm = FakeLLM(loop + [message(content="Washout removes carryover [S1].")])
    result = run_agent("washout?", retriever, llm)
    assert llm.calls[-1]["tool_choice"] == "none"
    assert result.answer.startswith("Washout") and result.error is None


def test_agent_reports_tool_errors_back_to_model(retriever):
    llm = FakeLLM([
        message(tool_calls=[tool_call("t1", "continuous_sample_size", {"design": "parallel", "hypothesis": "noninferiority",
                                                                        "delta": 0, "sd": 1})]),
        message(content="I need a margin."),
    ])
    result = run_agent("NI sample size", retriever, llm)
    tool_msg = [m for m in llm.calls[1]["messages"] if m["role"] == "tool"][0]
    assert tool_msg["content"].startswith("ERROR")
    assert result.answer == "I need a margin."


def test_agent_surfaces_api_errors(retriever):
    class RateLimitError(Exception):
        pass

    result = run_agent("q", retriever, FakeLLM([RateLimitError("429 too many")]))
    assert result.error and "rate limit" in result.error.lower()


# ------------------------------------------------------------------ workflows
def test_compare_documents(retriever):
    llm = FakeLLM([message(content="**Overview** ... [S1] and [S2] differ.")])
    res = compare_documents("carryover and washout", ["Estimand_Guidance_2020.pdf", "Crossover_Book_2015"], retriever, llm)
    assert res.error is None and "[S1]" in res.body
    prompt = llm.calls[0]["messages"][1]["content"]
    assert "Document: Estimand Guidance 2020" in prompt and "Document: Crossover Book 2015" in prompt


def test_checklists_are_well_formed():
    lists = load_checklists()
    assert {"estimands_e9r1", "sap_core_e9", "multiple_endpoints", "external_control", "adaptive_design",
            "pro_endpoint", "master_protocol"} <= set(lists)
    for cl in lists.values():
        ids = [i["id"] for i in cl["items"]]
        assert len(ids) == len(set(ids))
        assert all({"title", "requirement", "query"} <= set(i) for i in cl["items"])


def test_check_section_parses_json(retriever):
    cl = load_checklists()["estimands_e9r1"]
    reply = {"summary": "Mostly complete [S1].", "items": [
        {"id": "E1", "status": "met", "evidence": "the estimand is", "gap": "", "suggestion": "", "citations": ["S1"]},
        {"id": "E3", "status": "Partial", "evidence": "", "gap": "no strategy for rescue [S1]", "suggestion": "Add ...",
         "citations": ["S1", "S99"]},
    ]}
    llm = FakeLLM([message(content="Here you go:\n" + json.dumps(reply))])
    res = check_section("The estimand is the difference in CLUE comfort ...", cl, retriever, llm)
    assert res.error is None
    by_id = {i["id"]: i for i in res.items}
    assert by_id["E1"]["status"] == "met" and by_id["E3"]["status"] == "partial"
    assert by_id["E3"]["citations"] == ["S1"]
    assert by_id["E2"]["status"] == "not_assessed"
    assert llm.calls[0]["json_mode"] is True
    md = record_to_markdown(res.to_record())
    assert "| E3 |" in md and "Partial" in md


# ------------------------------------------------------------------ exports
def test_exports(retriever):
    llm = FakeLLM([
        message(tool_calls=[tool_call("t1", "search_knowledge_bank", {"query": "washout"})]),
        message(content="## Summary\n- Washout matters [S1]\n\n| A | B |\n|---|---|\n| 1 | **2** |\n\n```\ncode\n```"),
    ])
    rec = run_agent("washout?", retriever, llm).to_record()
    md = record_to_markdown(rec, with_passages=True)
    assert "## Sources" in md and "[S1]" in md and "> " in md
    data = record_to_docx(rec)
    from docx import Document

    doc = Document(io.BytesIO(data))
    assert len(doc.tables) == 1
    assert any("Washout matters" in p.text for p in doc.paragraphs)

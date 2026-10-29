"""Export answers, comparisons and section checks to Markdown or Word (.docx)."""
from __future__ import annotations

import io
import re

STATUS_LABEL = {"met": "Met", "partial": "Partial", "missing": "Missing",
                "not_applicable": "N/A", "not_assessed": "Not assessed"}


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------


def _sources_md(record: dict, include_uncited: bool) -> list[str]:
    out = []
    srcs = [s for s in record.get("sources", []) if include_uncited or s.get("cited")]
    if srcs:
        out.append("\n## Sources\n")
        for s in srcs:
            sec = f" - {s['section']}" if s.get("section") else ""
            out.append(f"- **[{s['tag']}]** {s['label']}{sec}")
    calcs = [c for c in record.get("calcs", []) if include_uncited or c.get("cited")]
    if calcs:
        out.append("\n## Calculations\n")
        for c in calcs:
            out.append(f"**[{c['tag']}]**\n```\n{c['text']}\n```")
    return out


def _items_md(items: list[dict]) -> list[str]:
    if not items:
        return []
    out = ["\n## Checklist\n", "| # | Item | Status | Evidence | Gap / suggestion |", "|---|---|---|---|---|"]
    for it in items:
        gap = " ".join(x for x in (it.get("gap", ""), it.get("suggestion", "")) if x)
        cells = [it["id"], it["title"], STATUS_LABEL.get(it["status"], it["status"]),
                 (it.get("evidence") or "").replace("\n", " "), gap.replace("\n", " ")]
        out.append("| " + " | ".join(c.replace("|", "/") for c in cells) + " |")
    return out


def record_to_markdown(record: dict, include_uncited: bool = False, with_passages: bool = False) -> str:
    lines = [f"# {record.get('title', 'ClinStat output')}", ""]
    meta = f"_Generated {record.get('created', '')} with {record.get('model', '')} by ClinStat._"
    lines.append(meta)
    if record.get("type") == "answer" and record.get("question"):
        lines += ["", f"**Question:** {record['question']}"]
    lines += ["", record.get("body", "")]
    lines += _items_md(record.get("items", []))
    lines += _sources_md(record, include_uncited)
    if with_passages:
        srcs = [s for s in record.get("sources", []) if include_uncited or s.get("cited")]
        if srcs:
            lines.append("\n## Source passages\n")
            for s in srcs:
                lines.append(f"**[{s['tag']}] {s['label']}**\n\n> " + s["text"].replace("\n", "\n> ") + "\n")
    return "\n".join(lines).strip() + "\n"


def records_to_markdown(records: list[dict], title: str = "ClinStat session") -> str:
    parts = [f"# {title}\n"]
    for r in records:
        parts.append(record_to_markdown(r).replace("\n# ", "\n## ", 1).replace("# ", "## ", 1))
        parts.append("\n---\n")
    return "\n".join(parts)


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------

_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[SC]\d+\])")


def _add_inline(paragraph, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif re.fullmatch(r"\[[SC]\d+\]", part):
            run = paragraph.add_run(part)
            run.font.superscript = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _markdown_to_docx(doc, md: str) -> None:
    lines = md.splitlines()
    i = 0
    in_code = False
    code: list[str] = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code))
                run.font.name = "Consolas"
                code, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        stripped = line.strip()
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{2,}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(ncol):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _add_inline(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            doc.add_heading(m.group(2).strip("* "), level=min(len(m.group(1)), 3))
        elif re.match(r"^[-*+]\s+", stripped):
            _add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*+]\s+", "", stripped))
        elif re.match(r"^\d+[.)]\s+", stripped):
            _add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped))
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            _add_inline(p, stripped.lstrip("> "))
            from docx.shared import Inches

            p.paragraph_format.left_indent = Inches(0.3)
        elif stripped:
            _add_inline(doc.add_paragraph(), stripped)
        i += 1


def record_to_docx(record: dict, include_uncited: bool = False, with_passages: bool = True) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    md = record_to_markdown(record, include_uncited=include_uncited, with_passages=with_passages)
    _markdown_to_docx(doc, md)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def records_to_docx(records: list[dict], title: str = "ClinStat session") -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    for r in records:
        _markdown_to_docx(doc, record_to_markdown(r))
        doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
